"""
Generador del Informe Administrativo oficial en formato Word (.docx)
conforme a la estructura del RAP para el Instructor Anybal López (SENA ADSO).
"""
import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from build_docs_helper import set_cell_background, set_cell_margins, set_table_borders, add_callout_box

def build_docx_from_md(md_file_path, output_docx_path):
    with open(md_file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # 1. Márgenes estándar (1 pulgada / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Estilo base
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Colores corporativos
    COLOR_PRIMARY = RGBColor(0x0F, 0x17, 0x2A)     # Slate 900
    COLOR_SECONDARY = RGBColor(0x1E, 0x3A, 0x8A)   # Blue 900
    COLOR_MUTED = RGBColor(0x47, 0x55, 0x69)       # Slate 600
    HEX_HEADER_BG = "1E3A8A"                        # Azul institucional
    HEX_ALT_ROW = "F8FAFC"                          # Gris muy suave

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        return p

    def add_para(text, bold_prefix=None, bullet=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        if bullet:
            p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = p.add_run("▪ ")
            r_bullet.bold = True
            r_bullet.font.color.rgb = COLOR_SECONDARY

        if bold_prefix:
            r_bp = p.add_run(bold_prefix)
            r_bp.bold = True
            r_bp.font.color.rgb = COLOR_PRIMARY

        # Procesar negritas markdown **texto**
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)
        return p

    def render_table(table_rows):
        if not table_rows:
            return
        
        num_cols = len(table_rows[0])
        table = doc.add_table(rows=len(table_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_borders(table, color="CBD5E1", sz="4", val="single")

        for r_idx, row in enumerate(table_rows):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
                
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.space_before = Pt(1)
                cell_p.paragraph_format.space_after = Pt(1)
                cell_p.paragraph_format.line_spacing = 1.05

                clean_text = val.strip()

                if r_idx == 0:
                    set_cell_background(cell, HEX_HEADER_BG)
                    run = cell_p.add_run(clean_text)
                    run.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if r_idx % 2 == 1:
                        set_cell_background(cell, HEX_ALT_ROW)
                    else:
                        set_cell_background(cell, "FFFFFF")
                    
                    # Chequear negritas
                    parts = re.split(r'(\*\*.*?\*\*)', clean_text)
                    for part in parts:
                        if part.startswith('**') and part.endswith('**'):
                            r = cell_p.add_run(part[2:-2])
                            r.bold = True
                            r.font.size = Pt(9)
                        else:
                            r = cell_p.add_run(part.replace('<br>', '\n'))
                            r.font.size = Pt(9)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Parser simple de Markdown
    i = 0
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Bloque de código
        if line.startswith('```'):
            if in_code_block:
                # Termina el bloque
                in_code_block = False
                code_text = '\n'.join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                cell.width = Inches(6.5)
                set_cell_background(cell, "F1F5F9")
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                cp = cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(2)
                cp.paragraph_format.space_after = Pt(2)
                r = cp.add_run(code_text)
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
                code_lines = []
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Tablas
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            parts = [c.strip() for c in line.strip().split('|')[1:-1]]
            # Si es separador de encabezado (|---|---|), omitir
            if all(re.match(r'^:?-+:?$', p) for p in parts if p):
                i += 1
                continue
            if not in_table:
                in_table = True
                table_rows = [parts]
            else:
                table_rows.append(parts)
            i += 1
            continue
        else:
            if in_table:
                render_table(table_rows)
                in_table = False
                table_rows = []

        # Saltos o separadores
        if line.strip() == '---':
            i += 1
            continue

        # Encabezados
        if line.startswith('# '):
            add_h1(line[2:].strip())
        elif line.startswith('## '):
            add_h2(line[3:].strip())
        elif line.startswith('### '):
            add_h3(line[4:].strip())
        elif line.startswith('#### '):
            add_h3(line[5:].strip())
        elif line.startswith('> [!NOTE]') or line.startswith('> [!IMPORTANT]') or line.startswith('> [!WARNING]'):
            tag = "NOTA INSTITUCIONAL"
            if "IMPORTANT" in line:
                tag = "IMPORTANTE"
            elif "WARNING" in line:
                tag = "ADVERTENCIA TÉCNICA"
            # Tomar líneas siguientes de blockquote
            bq_lines = []
            i += 1
            while i < len(lines) and lines[i].startswith('>'):
                bq_lines.append(lines[i].lstrip('> ').strip())
                i += 1
            add_callout_box(doc, ' '.join(bq_lines), title=tag)
            continue
        elif line.startswith('- ') or line.startswith('* '):
            add_para(line[2:].strip(), bullet=True)
        elif line.strip() == '':
            pass
        else:
            add_para(line.strip())

        i += 1

    if in_table:
        render_table(table_rows)

    doc.save(output_docx_path)
    print(f"Documento Word generado exitosamente en: {output_docx_path}")

if __name__ == '__main__':
    md_in = r"c:\Academix\INFORME_ADMINISTRATIVO_ACADEMIX.md"
    docx_out = r"c:\Academix\INFORME_ADMINISTRATIVO_ACADEMIX.docx"
    build_docx_from_md(md_in, docx_out)
