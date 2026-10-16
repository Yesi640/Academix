#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
generar_manual_despliegue_docx.py
Genera el documento Word (DOCX) ejecutivo del Manual de Despliegue y Acceso Web de ACADEMIX.
"""

import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_manual_docx():
    doc = Document()

    # Configuración de márgenes estándar (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Estilos de encabezado
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("ACADEMIX")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 76, 129)  # Azul Institucional

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("MANUAL DE DESPLIEGUE, CONFIGURACIÓN Y ACCESO WEB OFICIAL\nSistema Integral de Gestión Académica y Control Escolar")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(80, 80, 80)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_date = date_p.add_run("Fecha de Despliegue: Octubre 2026 | Estado: OPERATIVO EN VIVO (LIVE)")
    run_date.font.name = 'Arial'
    run_date.font.size = Pt(10)
    run_date.font.bold = True
    run_date.font.color.rgb = RGBColor(46, 125, 50)  # Verde Éxito

    doc.add_paragraph()

    # SECCION 1: ENLACES
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. ENLACES OFICIALES DEL SISTEMA EN PRODUCCIÓN")
    run_h1.font.name = 'Arial'
    run_h1.font.color.rgb = RGBColor(15, 76, 129)

    doc.add_paragraph("La aplicación se encuentra desplegada y funcionando en la nube de alta disponibilidad de Render, con certificado de seguridad SSL (HTTPS) activo:")

    table_links = doc.add_table(rows=5, cols=2)
    table_links.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Destino / Módulo", "URL de Acceso Directo"]
    for i, h in enumerate(headers):
        cell = table_links.cell(0, i)
        cell.text = h
        set_cell_background(cell, "0F4C81")
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    links_data = [
        ("Portal Web Institucional", "https://academix-q2g4.onrender.com"),
        ("Inicio de Sesión (Login)", "https://academix-q2g4.onrender.com/accounts/login/"),
        ("Registro Institucional", "https://academix-q2g4.onrender.com/accounts/register/"),
        ("Panel de Superadministración", "https://academix-q2g4.onrender.com/admin/"),
    ]

    for row_idx, (dest, url) in enumerate(links_data, start=1):
        c0 = table_links.cell(row_idx, 0)
        c1 = table_links.cell(row_idx, 1)
        c0.text = dest
        c1.text = url
        set_cell_margins(c0, top=100, bottom=100, left=150, right=150)
        set_cell_margins(c1, top=100, bottom=100, left=150, right=150)
        c0.paragraphs[0].runs[0].font.name = 'Arial'
        c0.paragraphs[0].runs[0].font.bold = True
        c1.paragraphs[0].runs[0].font.name = 'Courier New'
        c1.paragraphs[0].runs[0].font.size = Pt(9.5)
        if row_idx % 2 == 0:
            set_cell_background(c0, "F4F6F9")
            set_cell_background(c1, "F4F6F9")

    doc.add_paragraph()

    # SECCION 2: CREDENCIALES
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. CREDENCIALES MAESTRAS DE PRODUCCIÓN")
    run_h2.font.name = 'Arial'
    run_h2.font.color.rgb = RGBColor(15, 76, 129)

    doc.add_paragraph("El proceso de despliegue inicializa automáticamente las siguientes cuentas con privilegios jerárquicos:")

    table_creds = doc.add_table(rows=3, cols=4)
    table_creds.alignment = WD_TABLE_ALIGNMENT.CENTER
    creds_headers = ["Rol", "Usuario", "Contraseña", "Alcance y Permisos"]
    for i, h in enumerate(creds_headers):
        cell = table_creds.cell(0, i)
        cell.text = h
        set_cell_background(cell, "0F4C81")
        p = cell.paragraphs[0]
        p.runs[0].font.name = 'Arial'
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)

    creds_data = [
        ("Superadministrador", "admin", "Academix2026!", "Control total del sistema, auditoría inmutable, base de datos y configuración institucional."),
        ("Rector(a) Institucional", "rector", "Academix2026!", "Gestión pedagógica, asignación docente, observación de convivencia (Ley 1620) y boletines."),
    ]

    for row_idx, (rol, usr, pwd, scope) in enumerate(creds_data, start=1):
        cells = [table_creds.cell(row_idx, i) for i in range(4)]
        cells[0].text = rol
        cells[1].text = usr
        cells[2].text = pwd
        cells[3].text = scope
        for i, c in enumerate(cells):
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)
            c.paragraphs[0].runs[0].font.name = 'Arial'
            if i in [1, 2]:
                c.paragraphs[0].runs[0].font.bold = True
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(15, 76, 129)
            if row_idx % 2 == 0:
                set_cell_background(c, "F4F6F9")

    doc.add_paragraph()

    # SECCION 3: RESUMEN TECNICO
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. MODIFICACIONES Y BLINDAJE TÉCNICO IMPLEMENTADO")
    run_h3.font.name = 'Arial'
    run_h3.font.color.rgb = RGBColor(15, 76, 129)

    doc.add_paragraph("A continuación se detallan los componentes técnicos configurados para lograr un despliegue sin fallas:")

    points = [
        ("Servidor WSGI Gunicorn:", "Integrado en requirements.txt y Procfile para servir peticiones concurrentes de múltiples usuarios en Linux."),
        ("Entrega de Estáticos con WhiteNoise:", "Configurado para servir CSS, JavaScript, fuentes y logotipos con compresión Gzip/Brotli sin requerir servidores adicionales."),
        ("Soporte para Bases de Datos en la Nube:", "Integración de dj-database-url y psycopg2-binary, permitiendo conectar PostgreSQL o MySQL de manera nativa."),
        ("Endurecimiento de Seguridad SSL / HTTPS:", "Activación de SECURE_SSL_REDIRECT, cookies seguras (SESSION_COOKIE_SECURE, CSRF_COOKIE_SECURE), cabecera SECURE_PROXY_SSL_HEADER, HSTS de 1 año y protección anti-clickjacking (X-Frame-Options: DENY)."),
        ("Auditoría Oficial Aprobada:", "Ejecución de 'python manage.py check --deploy' con 0 advertencias y 0 errores reportados."),
        ("Automatización de Despliegue (build.sh):", "Compilación de 5 fases que actualiza pip, instala paquetes, recopila archivos estáticos, aplica migraciones e inicializa los datos maestros."),
        ("Inicialización Maestra (setup_production_init.py):", "Verificación y creación del Año Escolar 2026-2027, catálogo oficial de grados de Transición a 11°, y los 4 periodos académicos bimestrales."),
        ("Pruebas Unitarias al 100%:", "Corrección de servicios de auditoría y calificaciones bloqueadas; ejecución exitosa de pruebas con resultado OK."),
    ]

    for title, desc in points:
        p = doc.add_paragraph(style='List Bullet')
        r_t = p.add_run(title + " ")
        r_t.font.name = 'Arial'
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(15, 76, 129)
        r_d = p.add_run(desc)
        r_d.font.name = 'Arial'

    doc.add_paragraph()

    # SECCION 4: ACTUALIZACIONES FUTURAS
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. CÓMO ACTUALIZAR EL SISTEMA EN EL FUTURO")
    run_h4.font.name = 'Arial'
    run_h4.font.color.rgb = RGBColor(15, 76, 129)

    doc.add_paragraph("Render cuenta con integración continua (CI/CD) vinculada al repositorio oficial de GitHub. Para publicar futuras mejoras:")
    
    code_p = doc.add_paragraph()
    set_cell_background(table_links.cell(0,0), "0F4C81") # dummy
    r_code = code_p.add_run("git add .\ngit commit -m \"Descripción de la mejora\"\ngit push origin main")
    r_code.font.name = 'Courier New'
    r_code.font.size = Pt(10)
    r_code.font.bold = True
    r_code.font.color.rgb = RGBColor(33, 33, 33)

    doc.add_paragraph("Render detectará el push de inmediato, ejecutará build.sh y actualizará la web en 2 a 3 minutos de forma 100% transparente.")

    # Pie de página final
    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = footer_p.add_run("ACADEMIX © 2026 - Plataforma de Gestión y Auditoría Escolar Certificada")
    r_foot.font.name = 'Arial'
    r_foot.font.size = Pt(9)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(120, 120, 120)

    output_path = Path(r"c:\Academix\MANUAL_DESPLIEGUE_Y_ACCESO_WEB_ACADEMIX.docx")
    doc.save(output_path)
    print(f"[OK] Documento Word generado exitosamente: {output_path}")

if __name__ == '__main__':
    create_manual_docx()
